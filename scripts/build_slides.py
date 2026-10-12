"""One editable Markdown source -> editable PPTX and matching vector-text PDF."""
from __future__ import annotations
import re
import hashlib
from pathlib import Path
from slide_assets import build_assets
from common import ROOT, write_json

W, H = 13.333333, 7.5
INK, TEAL, GOLD, PAPER, WHITE = "131D2A", "007E87", "CFB991", "F4F3EF", "FFFFFF"
BLUE, ORANGE, MUTED, LINE = "345FB5", "B65131", "627081", "D9DFE3"


def parse():
    text = (ROOT/"slides/workshop.md").read_text(encoding="utf-8")
    slides = []
    for part in re.split(r"^## ", text, flags=re.M)[1:]:
        title, rest = part.split("\n", 1)
        number, title = title.split(" · ", 1)
        before, notes = rest.split("Notes:\n", 1)
        meta, bullets = {}, []
        for line in before.splitlines():
            if line.startswith("- "):
                bullets.append(line[2:])
            elif ": " in line:
                key, value = line.split(": ", 1)
                meta[key.lower()] = value
        slides.append({"number": number, "title": title, "bullets": bullets,
                       "notes": notes.strip(), **meta})
    previous = 600
    previous_slot = None
    main_slides = [s for s in slides if s["number"].isdigit()]
    breaks = []
    for slide in main_slides:
        a, b = slide["time"].split("–")
        minutes = lambda x: int(x.split(":")[0])*60+int(x.split(":")[1])
        slot = (a, b, slide.get("time group"))
        if slide.get("time group") and slot == previous_slot:
            continue  # Several app slides share one five-minute showcase slot.
        assert minutes(a) == previous, slide
        assert minutes(a) % 5 == 0 and minutes(b) % 5 == 0, slide
        assert minutes(b) > minutes(a), slide
        if slide["layout"] == "break":
            breaks.append(minutes(b) - minutes(a))
        previous = minutes(b)
        previous_slot = slot
    assert previous == 780 and len(main_slides) == 30 and len(slides) == 30
    assert breaks == [10, 10], "Keep two full ten-minute breaks"
    return slides


def main():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from PIL import Image
    from matplotlib.font_manager import findfont, FontProperties
    build_assets()
    slides = parse()
    out = ROOT/"slides"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = "Hands-On with Rumoca"
    prs.core_properties.author = "James Goppert · Micah Condie · Purdue University"
    prs.core_properties.subject = "American Modelica Conference · Georgia Tech · October 12, 2026"
    for name, family, weight in [("Workshop", "DejaVu Sans", "normal"), ("WorkshopBold", "DejaVu Sans", "bold"), ("WorkshopMono", "DejaVu Sans Mono", "normal")]:
        pdfmetrics.registerFont(TTFont(name, findfont(FontProperties(family=family, weight=weight))))
    pdf = canvas.Canvas(str(out/"workshop.pdf"), pagesize=(W*72, H*72))
    pdf.setTitle(prs.core_properties.title); pdf.setAuthor(prs.core_properties.author)
    audit, hyperlinks = [], []
    for spec in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(PAPER)
        pdf.setFillColor("#"+PAPER); pdf.rect(0, 0, W*72, H*72, fill=1, stroke=0)
        def rect(x, y, w, h, color, rounded=False):
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                           Inches(x), Inches(y), Inches(w), Inches(h))
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)
            shape.line.fill.background()
            style = shape._element.find(qn("p:style"))
            if style is not None: shape._element.remove(style)
            pdf.setFillColor("#"+color)
            if rounded:
                shape.adjustments[0] = .08
                pdf.roundRect(x*72, (H-y-h)*72, w*72, h*72, 8, fill=1, stroke=0)
            else:
                pdf.rect(x*72, (H-y-h)*72, w*72, h*72, fill=1, stroke=0)
        def text(value, x, y, w, size=28, color=INK, bold=False, mono=False, max_lines=None):
            # Wrap once using actual font metrics; both formats use identical lines.
            pdf_font = "WorkshopMono" if mono else "WorkshopBold" if bold else "Workshop"
            lines = []
            for paragraph in value.split("\n"):
                current = ""
                words = []
                for word in paragraph.split():
                    # File names and other long tokens still have to fit their panel.
                    while pdfmetrics.stringWidth(word, pdf_font, size) > w*72:
                        n = 1
                        while n < len(word) and pdfmetrics.stringWidth(word[:n+1], pdf_font, size) <= w*72:
                            n += 1
                        words.append(word[:n]); word = word[n:]
                    words.append(word)
                for word in words:
                    candidate = (current+" "+word).strip()
                    if pdfmetrics.stringWidth(candidate, pdf_font, size) > w*72 and current:
                        lines.append(current); current = word
                    else:
                        current = candidate
                lines.append(current)
            assert all(pdfmetrics.stringWidth(line, pdf_font, size) <= w*72+.1 for line in lines), (spec["number"], "line exceeds panel")
            if max_lines is not None:
                assert len(lines) <= max_lines, (spec["number"], value, lines)
            height = len(lines)*size*1.24/72 + .06
            assert y+height <= H-.13, (spec["number"], "text overflow", value)
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w+.06), Inches(height+.08))
            tf = shape.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = False
            for idx, line in enumerate(lines):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = line
                p.font.name = "DejaVu Sans Mono" if mono else "DejaVu Sans"
                p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
                p.space_before = p.space_after = Pt(0); p.line_spacing = 1.24
            pdf.setFont(pdf_font, size); pdf.setFillColor("#"+color)
            for i, line in enumerate(lines):
                pdf.drawString(x*72, (H-y)*72-size-i*size*1.24, line)
            audit.append({"slide": spec["number"], "text": value, "font_pt": size, "lines": len(lines), "x": x, "y": y, "width": w, "height": height})
            return height
        def picture(path, x, y, w, h):
            path = (out/path).resolve()
            with Image.open(path) as im: iw, ih = im.size
            scale = min(w/iw, h/ih)
            aw, ah = iw*scale, ih*scale
            xx, yy = x+(w-aw)/2, y+(h-ah)/2
            slide.shapes.add_picture(str(path), Inches(xx), Inches(yy), width=Inches(aw), height=Inches(ah))
            pdf.drawImage(str(path), xx*72, (H-yy-ah)*72, aw*72, ah*72, mask="auto")
        def movie(path, poster, x, y, w, h):
            # Embedded MP4 in PPTX; the same poster remains readable in PDF.
            with Image.open(out/poster) as im:
                iw, ih = im.size
            scale = min(w/iw, h/ih)
            aw, ah = iw*scale, ih*scale
            xx, yy = x+(w-aw)/2, y+(h-ah)/2
            slide.shapes.add_movie(str(out/path), Inches(xx), Inches(yy),
                                   Inches(aw), Inches(ah),
                                   poster_frame_image=str(out/poster), mime_type="video/mp4")
            pdf.drawImage(str(out/poster), xx*72, (H-yy-ah)*72, aw*72, ah*72, mask="auto")
            pdf.linkURL(spec["video link"], (xx*72, (H-yy-ah)*72,
                        (xx+aw)*72, (H-yy)*72), relative=0, thickness=0)
        # Each phase has one accent, reused in navigation, prompts and plots.
        num = int(spec["number"]) if spec["number"].isdigit() else 99
        phase = "MODEL" if num <= 14 else "FIT" if num <= 21 else "OPTIMIZE" if num <= 23 else "EXPLORE"
        accent = {"MODEL": TEAL, "FIT": BLUE, "OPTIMIZE": ORANGE, "EXPLORE": TEAL}[phase]
        layout = spec["layout"]
        dark = layout in {"title", "break", "race"} or num in {2, 29}
        bg, fg = (INK, WHITE) if dark else (PAPER, INK)
        rect(0, 0, W, H, bg)
        rect(0, 0, W, .065, GOLD)
        def rule(x, y, w, color=LINE):
            rect(x, y, w, .015, color)
        def label(value, x, y, w=10, color=accent):
            text(value.upper(), x, y, w, 11, color, bold=True, max_lines=1)
        def heading():
            label(spec.get("badge", ""), .6, .28, 11.1, GOLD if dark else accent)
            text(spec["time"], 11.5, .28, 1.25, 11, GOLD if dark else MUTED, max_lines=1)
            size = min(34, 12.1*72/pdfmetrics.stringWidth(spec["title"], "WorkshopBold", 34)*34)
            text(spec["title"], .6, .82, 12.1, size, fg, bold=True, max_lines=1)
        def stagebar(y=6.97):
            for i, (name, color) in enumerate([("MODEL", TEAL), ("FIT", BLUE), ("OPTIMIZE", ORANGE)]):
                x = 9.2+i*1.07
                rect(x, y, .13, .13, color)
                text(name, x+.2, y-.015, .84, 7.5, GOLD if dark else MUTED, bold=phase == name, max_lines=1)
        def steps(bullets, x, y, w, size=23, gap=.32):
            for i, bullet in enumerate(bullets):
                rect(x, y+.025, .37, .37, accent, rounded=True)
                text(str(i+1), x+.105, y+.055, .2, 14, WHITE, bold=True, max_lines=1)
                y += text(bullet, x+.55, y, w-.55, size, fg, max_lines=4)+gap
            return y
        if layout == "title":
            label(spec.get("badge", ""), .65, .35, 11, GOLD)
            text(spec["time"], 11.35, .35, 1.35, 11, GOLD, max_lines=1)
            if num == 1:
                label("ONE MODEL. SIMULATE. FIT. OPTIMIZE.", .7, 1.1, 9, GOLD)
                text("Hands-On", .65, 1.6, 9.3, 59, WHITE, bold=True, max_lines=1)
                text("with Rumoca", .65, 2.52, 9.8, 64, GOLD, bold=True, max_lines=1)
                text(spec["presenter 1"], .7, 3.73, 9.6, 23, WHITE, bold=True, max_lines=1)
                text(spec["role 1"], .7, 4.13, 9.45, 17, "D1D7DF", max_lines=1)
                text(spec["presenter 2"], .7, 4.53, 9.6, 23, WHITE, bold=True, max_lines=1)
                text(spec["role 2"], .7, 4.93, 9.45, 17, "D1D7DF", max_lines=1)
                text("Purdue University · School of Aeronautics and Astronautics", .7, 5.33, 9.45, 15, GOLD, max_lines=1)
                for i, (name, sub, color) in enumerate([("MODEL", "Declare the physics", TEAL), ("FIT", "Learn from data", BLUE), ("OPTIMIZE", "Plan the motion", ORANGE)]):
                    x = .7+i*3.12
                    rect(x, 5.68, 2.9, .7, color, rounded=True)
                    text(name, x+.16, 5.78, 2.5, 15, WHITE, bold=True, max_lines=1)
                    text(sub, x+.16, 6.09, 2.55, 10, WHITE, max_lines=1)
                text(spec["bullets"][3], .7, 6.45, 9.4, 10, GOLD, max_lines=1)
                picture(spec["purdue logo"], 10.43, 1.19, 1.99, .33)
                picture(spec["penn state logo"], 10.28, 1.62, 1.35, .39)
                picture(spec["nxp logo"], 11.85, 1.67, .68, .29)
                picture(spec["photo"], 10.2, 2.2, 2.45, 1.27)
                rect(10.2, 3.65, 2.45, 2.65, WHITE, rounded=True)
                picture(spec["visual"], 10.37, 3.82, 2.11, 2.11)
                text("CREATE CORE CODESPACE", 10.38, 6.02, 2.1, 8.5, INK, bold=True, max_lines=1)
            else:
                text("Bring your model.", .65, 1.6, 9.5, 48, WHITE, bold=True, max_lines=1)
                text("Join the Rumoca research community.", .65, 2.42, 12.0, 31, GOLD, bold=True, max_lines=1)
                steps(spec["bullets"], .7, 3.35, 11.9, size=22, gap=.23)
                email = spec["email"]
                shape = text(email, .7, 5.73, 11.9, 32, WHITE, bold=True, max_lines=1)
                # The address is also a clickable contact, in both formats.
                slide.shapes[-1].click_action.hyperlink.address = "mailto:"+email
                pdf.linkURL("mailto:"+email, (.7*72, (H-6.3)*72, 8.2*72, (H-5.7)*72), relative=0, thickness=0)
                hyperlinks.append({"slide":spec["number"], "label":email, "url":"mailto:"+email,
                                   "x":.7, "y":5.7, "width":7.5, "height":.6})
                text("James Goppert · Micah Condie", .7, 6.29, 9.5, 17, GOLD, max_lines=1)
        elif layout == "break":
            label("PAUSE / RECHARGE", .65, .4, 10, GOLD)
            text("Back at", .7, 1.35, 8, 35, WHITE, max_lines=1)
            text(spec["time"].split("–")[1], .62, 2.02, 8.3, 102, GOLD, bold=True, max_lines=1)
            rect(9.4, 1.3, 3.3, 4.85, "203044", rounded=True)
            label("KEEP YOUR PLACE", 9.65, 1.6, 2.8, GOLD)
            y=2.18
            for bullet in spec["bullets"]:
                y += text(bullet, 9.65, y, 2.75, 19, WHITE, max_lines=4)+.35
            text("10 MINUTES", .7, 4.35, 8.2, 20, WHITE, bold=True, max_lines=1)
            text("Stretch. Ask a question. Keep your editor open.", .7, 4.98, 8, 21, "D1D7DF", max_lines=2)
        elif layout == "race":
            heading()
            text(spec["bullets"][1], .7, 1.73, 9.05, 24, GOLD, bold=True, max_lines=1)
            movie(spec["movie"], spec["poster"], .7, 2.31, 9.05, 3.89)
            text("PURT FPV chase · September 2025 · click video to play", .7, 6.29, 9.05, 14, "D1D7DF", max_lines=1)
            rect(10.14, 2.39, 2.54, 3.54, WHITE, rounded=True)
            picture(spec["visual"], 10.33, 2.6, 2.16, 2.16)
            text("FOLLOW THE\nNEXT COMPETITION", 10.36, 5.05, 2.05, 12, INK, bold=True, max_lines=2)
        elif layout == "research":
            heading()
            text("A unified, declarative foundation for cyberphysical systems.", .7, 1.55, 11.9, 24, INK, bold=True, max_lines=1)
            for i, bullet in enumerate(spec["bullets"]):
                name, explanation = bullet.split(": ", 1)
                x, y = .65+(i%2)*6.17, 2.25+(i//2)*1.9
                color=[TEAL, BLUE, ORANGE, TEAL][i]
                rect(x, y, 5.82, 1.68, WHITE, rounded=True)
                rect(x, y, .07, 1.68, color)
                name_size = min(21, 5.22*72/pdfmetrics.stringWidth(name, "WorkshopBold", 21)*20.5)
                text(name, x+.27, y+.2, 5.22, name_size, color, bold=True, max_lines=1)
                text(explanation, x+.27, y+.8, 5.22, 18, INK, max_lines=2)
            text("Cerebri RDD2: Modelica → eFMI C → Zephyr RTOS → real drone",
                 .7, 6.1, 11.9, 19, TEAL, bold=True, max_lines=1)
        elif num in {2, 29}:
            heading()
            if num == 2:
                label("YOUR QUESTION. YOUR EXPERIMENT.", .7, 1.8, 8.6, GOLD)
                text("Make something you care about.", .65, 2.27, 8.2, 31, WHITE, bold=True, max_lines=1)
                label("EXAMPLE AGENT PROMPT", .7, 3.13, 8.2, GOLD)
                text(spec["prompt"], .7, 3.54, 8.2, 23, WHITE, max_lines=5)
                text("Choose any problem you care about.", .7, 5.3, 8.2, 20, GOLD, max_lines=1)
                text("Demo: agent/tool + elapsed time + one check", .7, 5.84, 8.2, 17, "D1D7DF", max_lines=1)
                rect(9.15, 1.78, 3.55, 4.57, "203044", rounded=True)
                label("THE PRIZE", 9.43, 2.08, 3, GOLD)
                text("FREE", 9.43, 2.66, 3, 41, GOLD, bold=True, max_lines=1)
                text("CogniPilot\nstocking cap", 9.43, 3.43, 2.95, 24, WHITE, bold=True, max_lines=2)
                rule(9.43, 4.6, 2.95, GOLD)
                text("Short demos\nAudience hand vote", 9.43, 4.94, 2.95, 19, WHITE, max_lines=2)
            else:
                for i, (big, small, color) in enumerate([("5", "MAX. PRESENTERS", TEAL), ("45 s", "EACH DEMO", BLUE), ("1", "VOTE PER PERSON", ORANGE)]):
                    x=.65+i*4.19
                    rect(x, 1.75, 3.97, 2.32, color, rounded=True)
                    text(big, x+.25, 1.94, 3.45, 60, WHITE, bold=True, max_lines=1)
                    text(small, x+.25, 3.43, 3.45, 13, WHITE, bold=True, max_lines=1)
                text("Show your change + one checked result", .7, 4.5, 11.9, 29, WHITE, bold=True, max_lines=1)
                text("Name your tool. Share elapsed work time.", .7, 5.14, 11.9, 23, WHITE, max_lines=1)
                text(spec["bullets"][2], .7, 5.77, 11.9, 23, GOLD, bold=True, max_lines=1)
        else:
            heading()
            if layout == "cards":
                for i, bullet in enumerate(spec["bullets"]):
                    x, y = .65+(i%2)*6.17, 1.75+(i//2)*2.15
                    color = [TEAL, BLUE, ORANGE, TEAL][i]
                    rect(x, y, 5.82, 1.95, WHITE, rounded=True)
                    text(f"0{i+1}", x+.25, y+.23, .72, 28, color, bold=True, max_lines=1)
                    text(bullet, x+1.13, y+.3, 4.4, 24, INK, max_lines=3)
                    rect(x+.25, y+1.62, .48, .055, color)
            elif layout == "image":
                image_h = 4.22 if "command" in spec else 4.45
                rect(4.82, 1.7, 7.86, image_h, WHITE, rounded=True)
                picture(spec["visual"], 4.94, 1.83, 7.62, image_h-.26)
                label("TRY THIS" if num in {5,7,11,26,27,28} else "READ THE MODEL" if num == 12 else "THE IDEA", .7, 1.86, 4)
                steps(spec["bullets"], .7, 2.35, 3.83, size=22, gap=.38)
            elif layout == "code":
                rect(.65, 1.73, 6.55, 4.7, INK, rounded=True)
                text(spec["code title"], .9, 1.95, 6.05, 14, GOLD, bold=True, max_lines=1)
                for i, line in enumerate(spec["code"].split(r"\n")):
                    if line:
                        indentation = len(line) - len(line.lstrip())
                        text(line.lstrip(), .9+indentation*.12, 2.43+i*.28,
                             6.05-indentation*.12, 16, WHITE, mono=True, max_lines=1)
                steps(spec["bullets"], 7.5, 2.04, 5.1, size=20, gap=.32)
            elif layout == "fit":
                rect(.65, 1.7, 12.03, 2.92, WHITE, rounded=True)
                picture(spec["visual"], .8, 1.78, 11.73, 2.76)
                label("PHYSICAL COEFFICIENTS", .8, 4.83, 5.7)
                text("kₗ: linear drag [N s/m]\nk_q: quadratic drag [N s²/m²]\nk_ref: one corresponding SI unit", .8, 5.12, 5.75, 18, INK, max_lines=3)
                label("WHAT THE OPTIMIZER CHANGES", 7.0, 4.83, 5.5)
                text("a, b: dimensionless log parameters\ne ≈ 2.718; eᵃ and eᵇ keep drag positive", 7.0, 5.12, 5.5, 18, INK, max_lines=3)
                text("ŷ: prediction · y: measurement · σ: noise standard deviation · L: mean squared scaled error",
                     .8, 6.18, 11.75, 16, MUTED, max_lines=1)
            elif layout == "wide":
                if spec["bullets"]:
                    rect(.65, 1.65, .045, .54, accent)
                    text(spec["bullets"][0], .85, 1.71, 11.6, 21, INK, bold=True, max_lines=2)
                py = 2.37 if spec["bullets"] else 1.7
                ph = (5.86 if "command" in spec else 6.21)-py
                rect(.65, py, 12.03, ph, WHITE, rounded=True)
                picture(spec["visual"], .81, py+.08, 11.71, ph-.16)
            elif layout == "firmware":
                label("PHYSICS + ACTUAL FIRMWARE", .7, 1.79, 4)
                steps(spec["bullets"], .7, 2.2, 3.83, size=19, gap=.28)
                rect(4.82, 1.7, 7.86, 4.22, WHITE, rounded=True)
                movie(spec["movie"], spec["visual"], 4.94, 1.85, 7.62, 3.82)
                text("Recorded research demo · separate from today’s 3 trials", 4.98, 5.75, 7.53, 12, MUTED, max_lines=1)
            elif layout == "bridge":
                text(spec["bullets"][0], .7, 1.75, 11.95, 22, INK, bold=True, max_lines=1)
                rect(.65, 2.41, 7.83, 3.78, WHITE, rounded=True)
                picture(spec["visual"], .8, 2.58, 7.53, 3.44)
                picture(spec["photo"], 8.77, 2.41, 3.91, 2.48)
                text("Digital twins →\nmeasured flight at PURT", 8.77, 5.16, 3.91, 20, TEAL, bold=True, max_lines=2)
            elif layout == "statement":
                if num == 3:
                    text(spec["bullets"][0], .7, 1.85, 11.9, 32, INK, bold=True, max_lines=2)
                    labels = [("SIMULATE", "What happens?", TEAL), ("IDENTIFY", "What fits the data?", BLUE), ("OPTIMIZE", "What should we do?", ORANGE)]
                    for i, (name, question, color) in enumerate(labels):
                        x=.65+i*4.19
                        rect(x, 3.22, 3.97, 1.67, color, rounded=True)
                        text(name, x+.24, 3.5, 3.5, 21, WHITE, bold=True, max_lines=1)
                        text(question, x+.24, 4.18, 3.5, 17, WHITE, max_lines=1)
                    text(spec["bullets"][2], .7, 5.43, 11.9, 24, INK, max_lines=2)
                else:
                    steps(spec["bullets"], .7, 2.08, 11.85, size=29, gap=.53)
            else:
                raise ValueError(layout)
        if "command" in spec:
            rect(.65, 6.08, 12.03, .48, INK, rounded=True)
            text("$ " + spec["command"], .85, 6.18, 11.7, 18, WHITE, mono=True, max_lines=1)
        cp = spec["checkpoint"]
        rect(.65, 6.68, .06, .25, GOLD if dark else accent)
        cp_size=min(18, 11.83*72/pdfmetrics.stringWidth(cp, "WorkshopBold", 18)*18)
        text(cp, .85, 6.67, 11.83, cp_size, GOLD if dark else accent, bold=True, max_lines=1)
        links = [entry.split(" | ", 1) for entry in spec.get("links", "").split(" ; ") if entry]
        if links:
            button_w = (11.72 - .12*(len(links)-1))/len(links)
            for i, (caption, url) in enumerate(links):
                assert url.startswith(("https://", "mailto:")), (spec["number"], url)
                x, y, h = .65+i*(button_w+.12), 7.035, .40
                first_shape = len(slide.shapes)
                rect(x, y, button_w, h, GOLD if dark else accent, rounded=True)
                label_text = caption+" ↗"
                label_size = min(13, (button_w-.24)*72/pdfmetrics.stringWidth(label_text, "WorkshopBold", 13)*13)
                text(label_text, x+.12, y+.04, button_w-.24, label_size,
                     INK if dark else WHITE, bold=True, max_lines=1)
                # Both the label and background are full clickable shapes.
                for shape in list(slide.shapes)[first_shape:]:
                    shape.click_action.hyperlink.address = url
                pdf.linkURL(url, (x*72, (H-y-h)*72, (x+button_w)*72, (H-y)*72),
                            relative=0, thickness=0)
                hyperlinks.append({"slide":spec["number"], "label":caption, "url":url,
                                   "x":x, "y":y, "width":button_w, "height":h})
            text(spec["number"], 12.5, 7.12, .35, 10, GOLD if dark else accent, bold=True, max_lines=1)
        else:
            rule(.65, 7.13, 12.03, "344153" if dark else LINE)
            text("PURDUE AEROSPACE  /  GEORGIA TECH  /  AMC 2026", .7, 7.17, 8.15, 8, "D1D7DF" if dark else MUTED, max_lines=1)
            stagebar(7.18)
            text(spec["number"], 12.43, 7.13, .3, 10, GOLD if dark else accent, bold=True, max_lines=1)
        slide.notes_slide.notes_text_frame.text = spec["time"]+"\n"+spec["notes"]
        pdf.showPage()
    prs.save(str(out/"workshop.pptx"))
    pdf.save()
    inputs = {"slides/workshop.md", "models/Quadrotor.mo", "data/noisy_flight.csv",
              "results/reference/fit_drag.json", "results/reference/trajectory.npz",
              "scripts/build_slides.py", "scripts/slide_assets.py", "pyproject.toml", "uv.lock"}
    media_keys = ["visual", "movie", "poster", "photo", "purdue logo", "penn state logo", "nxp logo"]
    inputs.update(str((out/s[key]).resolve().relative_to(ROOT))
                  for s in slides for key in media_keys if key in s)
    digest = lambda path: hashlib.sha256((ROOT/path).read_bytes()).hexdigest()
    write_json(out/"build-report.json", {"slides": len(slides), "main_slides": len([s for s in slides if s["number"].isdigit()]),
               "schedule_minutes": 180, "break_minutes": 20, "source": "slides/workshop.md",
               "input_sha256": {p: digest(p) for p in sorted(inputs)},
               "artifact_sha256": {p: digest(p) for p in ["slides/workshop.pptx", "slides/workshop.pdf"]},
               "rendered_text": audit, "hyperlinks": hyperlinks, "notes_in_pptx": True,
               "embedded_movies": [{"slide":s["number"], "source":"slides/"+s["movie"],
                                    "sha256":digest("slides/"+s["movie"])} for s in slides if "movie" in s]})
    print(f"Built {len(slides)} slides, PPTX + PDF, 180-minute schedule, speaker notes")


if __name__ == "__main__":
    main()
